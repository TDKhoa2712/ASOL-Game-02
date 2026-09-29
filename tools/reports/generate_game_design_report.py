"""Export the current CanDoKu GDD, never a parallel hard-coded design.

Requires python-docx. The Markdown files under GDD remain authoritative.
Generated Word documents require separate visual review before distribution.
"""
from pathlib import Path
import re
import sys

from docx import Document

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]


def build_game_design_document(output_path):
    document = Document()
    document.add_heading("CanDoKu — Game Design Document", 0)
    document.add_paragraph("Tìm kẹo bị đánh rơi trong vườn. Nguồn chuẩn: GDD/*.md; tiến độ: docs/STATUS.md.")
    sources = [REPOSITORY_ROOT / "GDD/README.md"]
    sources += sorted((REPOSITORY_ROOT / "GDD").glob("[0-9][0-9]-*.md"))
    for source in sources:
        document.add_page_break()
        document.add_paragraph(f"Nguồn: {source.relative_to(REPOSITORY_ROOT).as_posix()}")
        code = False
        for line in source.read_text(encoding="utf-8").splitlines():
            if line.startswith("```"):
                code = not code
                continue
            if code:
                document.add_paragraph(line, style="No Spacing")
            elif match := re.match(r"^(#{1,6}) (.+)", line):
                document.add_heading(match[2], min(len(match[1]), 4))
            elif line.strip():
                document.add_paragraph(line)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


if __name__ == "__main__":
    target = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else REPOSITORY_ROOT / "docs/reports/CanDoKu_GDD.docx"
    print(build_game_design_document(target))
